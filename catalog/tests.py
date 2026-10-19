from io import BytesIO
from collections import Counter
from tempfile import TemporaryDirectory
from unittest.mock import patch
from django.contrib.auth.models import Group, User
from django.contrib.auth.hashers import make_password
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.core.management.base import CommandError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook
from .importing import confirm, preview, headers, auto_mapping, set_preview_sample
from .models import Box, Customer, Destination, ImportJob, Integration, Part, Change, SheetColor
from .services import create_part, update_part

class CatalogTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        for role in ('Administrador', 'Editor', 'Consulta'):
            Group.objects.create(name=role)
        cls.editor = User.objects.create_user('editor', password='test-password')
        cls.editor.groups.add(Group.objects.get(name='Editor'))
        cls.reader = User.objects.create_user('reader', password='test-password')
        cls.reader.groups.add(Group.objects.get(name='Consulta'))
        cls.customer = Customer.objects.create(name='LENNOX 1')

    def test_create_duplicate_and_web_permissions(self):
        self.client.force_login(self.editor)
        response = self.client.post(reverse('catalog:create'), {'customer': self.customer.pk, 'part_number': '  ab-001  ', 'diameter': 'A - 3/8'})
        self.assertEqual(response.status_code, 302)
        part = Part.objects.get()
        self.assertEqual(part.part_number, 'ab-001')
        self.assertEqual(part.normalized_number, 'AB-001')
        self.assertEqual(Change.objects.count(), 1)
        response = self.client.post(reverse('catalog:create'), {'customer': self.customer.pk, 'part_number': 'AB-001'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Part.objects.count(), 1)
        self.client.force_login(self.reader)
        self.assertEqual(self.client.get(reverse('catalog:create')).status_code, 403)
        self.assertEqual(self.client.get(reverse('catalog:detail', args=[part.pk])).status_code, 200)

    def test_optimistic_edit_and_toggle(self):
        part = create_part({'customer': self.customer, 'part_number': 'P-1'}, self.editor)
        update_part(part.pk, {'diameter': 'A - 3/8'}, 1, self.editor)
        with self.assertRaises(ValidationError):
            update_part(part.pk, {'diameter': 'B'}, 1, self.editor)
        update_part(part.pk, {'active': False}, 2, self.editor)
        update_part(part.pk, {'active': True}, 3, self.editor)
        self.assertEqual(list(Change.objects.values_list('operation', flat=True)), ['create', 'update', 'deactivate', 'reactivate'])

    def make_job(self, rows):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Piezas'
        ws.append(['CLIENTE', 'NÚMERO DE PARTE', 'PARED'])
        for row in rows:
            ws.append(row)
        stream = BytesIO()
        wb.save(stream)
        return ImportJob.objects.create(actor=self.editor, filename='parts.xlsx', file=SimpleUploadedFile('parts.xlsx', stream.getvalue()), sheet='Piezas', mapping={'customer': 1, 'part_number': 2, 'wall': 3}, mode='both')

    def test_preview_confirm_repeat_and_conflict(self):
        job = self.make_job([['LENNOX 1', 'P-10', 0], ['LENNOX 1', 'P-11', 0.42]])
        rows = preview(job)
        self.assertEqual([row['status'] for row in rows], ['new', 'new'], rows)
        self.assertEqual(Part.objects.count(), 0)
        confirm(job.pk, self.editor)
        self.assertEqual(Part.objects.count(), 2)
        confirm(job.pk, self.editor)
        self.assertEqual(Part.objects.count(), 2)
        self.assertEqual(Part.objects.get(part_number='P-10').wall, 0)
        duplicate = self.make_job([['LENNOX 1', 'P-10', 1], ['LENNOX 1', 'P-10', 2]])
        self.assertEqual([r['status'] for r in preview(duplicate)], ['duplicate', 'duplicate'])
        with self.assertRaises(ValidationError):
            confirm(duplicate.pk, self.editor)

    def test_mapping_key_order_does_not_invalidate_preview(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Piezas'
        ws.append(['CLIENTE', 'DESTINO', 'NÚMERO DE PARTE', 'CAJA'])
        ws.append(['NEW CUSTOMER', 'NEW DESTINATION', 'P-12', 'NEW BOX'])
        stream = BytesIO()
        wb.save(stream)
        job = ImportJob.objects.create(
            actor=self.editor,
            filename='parts.xlsx',
            file=SimpleUploadedFile('parts.xlsx', stream.getvalue()),
            sheet='Piezas',
            mapping={'customer': 1, 'destination': 2, 'part_number': 3, 'box': 4},
            mode='both',
        )
        preview(job)
        job.result['rows'][0]['catalog_new'].reverse()
        job.save(update_fields=['result'])
        job.mapping = dict(reversed(list(job.mapping.items())))
        job.save(update_fields=['mapping'])
        confirm(job.pk, self.editor)
        part = Part.objects.get(part_number='P-12')
        self.assertEqual(part.customer.name, 'NEW CUSTOMER')
        self.assertEqual(part.destination.name, 'NEW DESTINATION')
        self.assertEqual(part.box.name, 'NEW BOX')

    def test_only_owner_can_delete_unconfirmed_import(self):
        with TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            job = self.make_job([['LENNOX 1', 'DELETE-1', 1]])
            preview(job)
            file_name = job.file.name
            url = reverse('catalog:import_delete', args=[job.pk])
            other_editor = User.objects.create_user('other-editor', password='test-password')
            other_editor.groups.add(Group.objects.get(name='Editor'))
            self.client.force_login(other_editor)
            self.assertEqual(self.client.post(url).status_code, 404)
            self.assertTrue(ImportJob.objects.filter(pk=job.pk).exists())

            self.client.force_login(self.editor)
            self.assertEqual(self.client.get(url).status_code, 405)
            self.assertEqual(self.client.post(url).status_code, 302)
            self.assertFalse(ImportJob.objects.filter(pk=job.pk).exists())
            self.assertFalse(job.file.storage.exists(file_name))
            self.assertFalse(Part.objects.filter(part_number='DELETE-1').exists())

            confirmed = self.make_job([['LENNOX 1', 'KEEP-1', 1]])
            preview(confirmed)
            confirm(confirmed.pk, self.editor)
            confirmed_url = reverse('catalog:import_delete', args=[confirmed.pk])
            self.assertEqual(self.client.post(confirmed_url).status_code, 302)
            self.assertTrue(ImportJob.objects.filter(pk=confirmed.pk).exists())
            self.assertTrue(Part.objects.filter(part_number='KEEP-1').exists())

    def test_import_stale_preview_rolls_back(self):
        part = create_part({'customer': self.customer, 'part_number': 'P-20'}, self.editor)
        job = self.make_job([['LENNOX 1', 'P-20', 1], ['LENNOX 1', 'P-21', 2]])
        preview(job)
        update_part(part.pk, {'wall': 3}, 1, self.editor)
        with self.assertRaises(ValidationError):
            confirm(job.pk, self.editor)
        self.assertFalse(Part.objects.filter(part_number='P-21').exists())
        self.assertEqual(Part.objects.get(pk=part.pk).wall, 3)

    def test_numeric_code_and_formula_are_row_errors(self):
        job = self.make_job([['LENNOX 1', 123, 1], ['LENNOX 1', 'SAFE-1', '=1+1']])
        rows = preview(job)
        self.assertEqual([row['status'] for row in rows], ['error', 'error'])
        with self.assertRaises(ValidationError):
            confirm(job.pk, self.editor)
        self.assertEqual(Part.objects.count(), 0)

    def test_reference_header_ambiguity(self):
        from pathlib import Path
        path = Path(__file__).resolve().parent.parent / 'UNIVERSO RAMOS.XLSX'
        if not path.exists():
            self.skipTest('Reference workbook is not present')
        job = ImportJob.objects.create(actor=self.editor, filename=path.name, file=SimpleUploadedFile(path.name, path.read_bytes()), sheet='Sheet1')
        _, names = headers(job, 'Sheet1')
        mapping = auto_mapping(names)
        self.assertEqual(mapping['part_number'], 4)
        self.assertEqual(mapping['customer'], 2)
        self.assertEqual(mapping['reference_customer'], 12)
        self.assertEqual(mapping['reference_plant'], 13)
        job.mapping = mapping
        job.save(update_fields=['mapping'])
        rows = preview(job)
        self.assertEqual(len(rows), 302)
        self.assertEqual(Counter(row['status'] for row in rows), {'new': 302}, [(row['row'], row['raw'], row['reason']) for row in rows if row['status'] != 'new'][:20])
        self.assertEqual(Part.objects.count(), 0)
        self.assertEqual(Customer.objects.count(), 1)
        self.assertEqual(Destination.objects.count(), 0)
        self.assertEqual(Box.objects.count(), 0)
        self.assertEqual(SheetColor.objects.count(), 0)
        self.client.force_login(self.editor)
        preview_page = self.client.get(reverse('catalog:import_preview', args=[job.pk]))
        self.assertContains(preview_page, 'CLIENTE (2)')
        self.assertContains(preview_page, 'PLANTA')
        self.assertNotContains(preview_page, 'CLIENTE no existe')
        confirm(job.pk, self.editor)
        self.assertEqual(Part.objects.count(), 302)
        self.assertEqual(Part.objects.filter(wall_note='NO').count(), 4)
        self.assertTrue(Part.objects.exclude(reference_customer='').exists())
        imported = Part.objects.get(part_number='82-105077-15')
        self.assertIsNone(imported.wall)
        self.assertEqual(imported.wall_note, 'NO')
        excel_page = self.client.get(reverse('catalog:list'), {'view': 'excel', 'q': imported.part_number})
        self.assertContains(excel_page, 'NO')
        self.assertContains(excel_page, imported.reference_customer)
        confirm(job.pk, self.editor)
        self.assertEqual(Part.objects.count(), 302)
        other = ImportJob.objects.create(actor=self.editor, filename=path.name, file=SimpleUploadedFile(path.name, path.read_bytes()), sheet='Sheet1 (2)')
        _, other_names = headers(other, other.sheet)
        other.mapping = auto_mapping(other_names)
        other.save(update_fields=['mapping'])
        other_rows = preview(other)
        self.assertGreaterEqual(Counter(row['status'] for row in other_rows)['duplicate'], 4)
        with self.assertRaises(ValidationError):
            confirm(other.pk, self.editor)

    def test_main_pages_render_and_actions_exist(self):
        self.client.force_login(self.editor)
        part = create_part({'customer': self.customer, 'part_number': 'VIS-1'}, self.editor)
        for url in [reverse('catalog:list'), reverse('catalog:detail', args=[part.pk]), reverse('catalog:create'), reverse('catalog:edit', args=[part.pk]), reverse('catalog:import_upload'), reverse('catalog:imports'), reverse('catalog:related', args=['clientes'])]:
            response = self.client.get(url)
            self.assertEqual(response.status_code, 200, url)
        self.assertContains(self.client.get(reverse('catalog:list')), 'Nuevo número de parte')
        excel_response = self.client.get(reverse('catalog:list'), {'view': 'excel'})
        self.assertContains(excel_response, 'CLIENTE (2)')
        self.assertContains(excel_response, 'PLANTA')
        job = self.make_job([['LENNOX 1', 'VIS-2', 1]])
        self.assertEqual(self.client.get(reverse('catalog:import_map', args=[job.pk])).status_code, 200)
        preview(job)
        self.assertEqual(self.client.get(reverse('catalog:import_preview', args=[job.pk])).status_code, 200)
        self.editor.groups.add(Group.objects.get(name='Administrador'))
        for url in [reverse('catalog:history'), reverse('catalog:users'), reverse('catalog:user_create'), reverse('catalog:integrations')]:
            self.assertEqual(self.client.get(url).status_code, 200, url)
        self.assertIn(b'VIS-1', self.client.get(reverse('catalog:export')).content)
        self.assertTrue(self.client.get(reverse('catalog:import_template')).content.startswith(b'PK'))

    def test_bootstrap_admin_once(self):
        with patch.dict('os.environ', {'BOOTSTRAP_ADMIN_HASH': make_password('test-only-password')}):
            call_command('bootstrap_admin', verbosity=0)
            admin = User.objects.get(username='administrator')
            self.assertTrue(admin.is_superuser)
            self.assertTrue(admin.check_password('test-only-password'))
            with self.assertRaises(CommandError):
                call_command('bootstrap_admin', verbosity=0)

    def test_api_auth_and_paged_changes(self):
        integration, token = Integration.create_key('MESA')
        self.assertEqual(self.client.get('/api/v1/parts/').status_code, 401)
        headers = {'HTTP_AUTHORIZATION': f'Bearer {token}'}
        part = create_part({'customer': self.customer, 'part_number': 'A-1'}, self.editor)
        update_part(part.pk, {'active': False}, 1, self.editor)
        first = self.client.get('/api/v1/changes/?limit=1', **headers).json()
        self.assertEqual(first['results'][0]['operation'], 'create')
        self.assertTrue(first['has_more'])
        second = self.client.get('/api/v1/changes/', {'limit': 1, 'cursor': first['next_cursor']}, **headers).json()
        self.assertEqual(second['results'][0]['operation'], 'deactivate')
        self.assertFalse(second['results'][0]['part']['active'])
        repeated = self.client.get('/api/v1/changes/?limit=1', **headers).json()
        self.assertEqual(repeated['results'][0]['sequence'], first['results'][0]['sequence'])
        self.assertEqual(self.client.get('/api/v1/parts/?page_size=1', **headers).json()['count'], 1)
        integration.active = False
        integration.save()
        self.assertEqual(self.client.get('/api/v1/parts/', **headers).status_code, 401)

    def test_sample_form_display_export_and_filters(self):
        self.client.force_login(self.editor)
        response = self.client.post(reverse('catalog:create'), {
            'customer': self.customer.pk, 'part_number': 'SAMPLE-1', 'is_sample': 'on',
        })
        self.assertEqual(response.status_code, 302)
        part = Part.objects.get(part_number='SAMPLE-1')
        self.assertTrue(part.is_sample)
        self.assertTrue(part.changes.get().after['is_sample'])
        self.assertContains(self.client.get(reverse('catalog:list')), '(MUESTRA)')
        self.assertContains(self.client.get(reverse('catalog:detail', args=[part.pk])), '(MUESTRA)')
        self.assertNotContains(self.client.get(reverse('catalog:list'), {'is_sample': 'false'}), 'SAMPLE-1')
        self.assertContains(self.client.get(reverse('catalog:export')), 'MUESTRA')
        self.client.post(reverse('catalog:edit', args=[part.pk]), {
            'customer': self.customer.pk, 'part_number': part.part_number, 'version': part.version,
        })
        part.refresh_from_db()
        self.assertFalse(part.is_sample)
        self.assertFalse(part.changes.last().after['is_sample'])

    def test_sample_import_preserves_and_changes_classification(self):
        part = create_part({'customer': self.customer, 'part_number': 'S-1', 'is_sample': True}, self.editor)
        job = self.make_job([['LENNOX 1', 'S-1', 1]])
        preview(job)
        confirm(job.pk, self.editor)
        part.refresh_from_db()
        self.assertTrue(part.is_sample)
        wb = Workbook()
        ws = wb.active
        ws.title = 'Piezas'
        ws.append(['CLIENTE', 'NÚMERO DE PARTE', 'MUESTRA'])
        ws.append(['LENNOX 1', 'S-1', 'No'])
        ws.append(['LENNOX 1', 'S-2', 'Sí'])
        ws.append(['LENNOX 1', 'S-3', False])
        stream = BytesIO()
        wb.save(stream)
        job = ImportJob.objects.create(actor=self.editor, filename='samples.xlsx',
            file=SimpleUploadedFile('samples.xlsx', stream.getvalue()), sheet='Piezas',
            mapping={'customer': 1, 'part_number': 2, 'is_sample': 3})
        rows = preview(job)
        self.assertEqual(rows[0]['before']['is_sample'], True)
        self.assertEqual(rows[0]['after']['is_sample'], False)
        self.assertTrue(rows[1]['is_sample'])
        confirm(job.pk, self.editor)
        self.assertFalse(Part.objects.get(part_number='S-1').is_sample)
        self.assertTrue(Part.objects.get(part_number='S-2').is_sample)
        self.assertFalse(Part.objects.get(part_number='S-3').is_sample)
        job = self.make_job([['LENNOX 1', 'S-2', 1]])
        job.mapping['is_sample'] = job.mapping.pop('wall')
        job.save()
        self.assertEqual(preview(job)[0]['status'], 'unchanged')
        # Blank sample values preserve the mark, unless clearing blanks is explicit.
        job = self.make_job([['LENNOX 1', 'S-2', None]])
        job.mapping['is_sample'] = job.mapping.pop('wall')
        job.save()
        self.assertEqual(preview(job)[0]['status'], 'unchanged')
        job.clear_blanks = True
        job.save()
        self.assertFalse(preview(job)[0]['after']['is_sample'])
        confirm(job.pk, self.editor)
        self.assertFalse(Part.objects.get(part_number='S-2').is_sample)

    def test_preview_filters_paginate_and_cannot_hide_blocking_errors(self):
        self.client.force_login(self.editor)
        job = self.make_job([['LENNOX 1', 'PREVIEW', 1]])
        rows = [{'row': i, 'status': 'unchanged', 'raw': {}, 'after': {}} for i in range(1, 62)]
        rows += [
            {'row': 62, 'status': 'new', 'raw': {'part_number': 'NEW'}, 'after': {}},
            {'row': 63, 'status': 'update', 'raw': {'part_number': 'UPDATED'}, 'after': {}},
            {'row': 64, 'status': 'error', 'raw': {}, 'after': {}, 'reason': 'Invalid'},
        ]
        job.result = {'rows': rows, 'counts': {'new': 1, 'update': 1, 'unchanged': 61, 'error': 1}}
        job.status = 'preview'
        job.save()
        url = reverse('catalog:import_preview', args=[job.pk])
        response = self.client.get(url, {'status': 'changed'})
        self.assertEqual([row['row'] for row in response.context['page']], [62, 63])
        self.assertTrue(response.context['blocked'])
        self.assertNotContains(response, '>Confirmar importación</button>')
        self.assertEqual(len(self.client.get(url, {'status': 'new'}).context['page']), 1)
        self.assertEqual(len(self.client.get(url, {'status': 'update'}).context['page']), 1)
        self.assertContains(self.client.get(url, {'status': 'skipped'}), 'No hay filas para este filtro')
        response = self.client.get(url, {'status': 'unchanged'})
        self.assertEqual(response.context['page'].paginator.count, 61)
        self.assertContains(response, 'status=unchanged&amp;page=2')

    def test_sample_api_filters_and_reclassification_events(self):
        _, token = Integration.create_key('Sample consumer')
        headers = {'HTTP_AUTHORIZATION': f'Bearer {token}'}
        part = create_part({'customer': self.customer, 'part_number': 'API-SAMPLE'}, self.editor)
        update_part(part.pk, {'is_sample': True}, 1, self.editor)
        self.assertEqual(self.client.get('/api/v1/parts/?is_sample=false', **headers).json()['count'], 0)
        self.assertTrue(self.client.get('/api/v1/parts/?is_sample=true', **headers).json()['results'][0]['is_sample'])
        self.assertEqual(self.client.get('/api/v1/parts/exact/?code=API-SAMPLE&is_sample=false', **headers).json()['count'], 0)
        self.assertEqual(self.client.get(f'/api/v1/parts/{part.pk}/?is_sample=false', **headers).status_code, 404)
        events = self.client.get('/api/v1/changes/', **headers).json()['results']
        self.assertFalse(events[0]['part']['is_sample'])
        self.assertTrue(events[1]['part']['is_sample'])
        update_part(part.pk, {'is_sample': False}, 2, self.editor)
        self.assertEqual(self.client.get('/api/v1/parts/?is_sample=false', **headers).json()['count'], 1)

    def test_consumer_sample_policy_and_reclassification(self):
        import json
        import runpy
        import sqlite3
        from pathlib import Path

        events = []
        parts = [{'id': 'sample-id', 'is_sample': True}, {'id': 'normal-id', 'is_sample': False}]

        def response(request, **kwargs):
            url = request.full_url
            if '/bootstrap/' in url:
                payload = {'changes_cursor': 'start'}
            elif '/parts/' in url:
                payload = {'results': parts, 'next': None}
            else:
                payload = {'results': events, 'next_cursor': 'next', 'has_more': False}
            return BytesIO(json.dumps(payload).encode())

        with TemporaryDirectory() as directory:
            cache = str(Path(directory) / 'cache.sqlite3')
            env = {'UR_BASE_URL': 'https://example.test', 'UR_API_KEY': 'test', 'UR_CACHE_DB': cache}

            def sync(include_samples):
                with patch.dict('os.environ', {**env, 'UR_INCLUDE_SAMPLES': include_samples}), patch('urllib.request.urlopen', side_effect=response):
                    result = runpy.run_path(str(Path(__file__).resolve().parent.parent / 'scripts/sample_consumer.py'))
                    result['db'].close()
                with sqlite3.connect(cache) as db:
                    return {row[0] for row in db.execute('SELECT central_uuid FROM parts')}

            self.assertEqual(sync('false'), {'normal-id'})

            self.assertEqual(sync('true'), {'normal-id', 'sample-id'})
            self.assertEqual(sync('false'), {'normal-id'})
            events[:] = [{'part_id': 'sample-id', 'part': {'id': 'sample-id', 'is_sample': False}}]
            self.assertEqual(sync('false'), {'normal-id', 'sample-id'})
            events[:] = [{'part_id': 'sample-id', 'part': {'id': 'sample-id', 'is_sample': True}}]
            self.assertEqual(sync('false'), {'normal-id'})

    def test_preview_sample_checkboxes_persist_and_confirm(self):
        self.client.force_login(self.editor)
        part = create_part({'customer': self.customer, 'part_number': 'EXISTING', 'wall': 1}, self.editor)
        job = self.make_job([['LENNOX 1', 'NEW', 1], ['LENNOX 1', 'EXISTING', 1]])
        preview(job)
        url = reverse('catalog:import_preview', args=[job.pk])
        self.assertContains(self.client.get(url), 'name="is_sample"', count=2)
        response = self.client.post(url + '?status=all', {
            'action': 'set_sample', 'row': 2, 'is_sample': 'on', 'expected_sample': 'false',
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response['Location'], url + '?status=all#sample-row-2')
        set_preview_sample(job.pk, self.editor, 3, True, False)
        job.refresh_from_db()
        self.assertEqual(job.sample_overrides, {'2': True, '3': True})
        self.assertEqual(job.result['counts'], {'new': 1, 'update': 1})
        self.assertEqual(Part.objects.count(), 1)
        part.refresh_from_db()
        self.assertFalse(part.is_sample)
        self.assertContains(self.client.get(url, {'status': 'update'}), 'checked')
        self.assertTrue(self.client.get(url, {'status': 'new'}).context['page'][0]['is_sample'])
        confirm(job.pk, self.editor)
        self.assertTrue(Part.objects.get(part_number='NEW').is_sample)
        part.refresh_from_db()
        self.assertTrue(part.is_sample)
        self.assertEqual(Change.objects.count(), 3)
        confirm(job.pk, self.editor)
        self.assertEqual(Change.objects.count(), 3)
        self.assertNotContains(self.client.get(url), 'name="action" value="set_sample"')

    def test_preview_can_uncheck_existing_and_excel_samples(self):
        part = create_part({'customer': self.customer, 'part_number': 'SAMPLE', 'is_sample': True}, self.editor)
        job = self.make_job([['LENNOX 1', 'SAMPLE', None], ['LENNOX 1', 'NEW-SAMPLE', 1]])
        job.mapping['is_sample'] = job.mapping.pop('wall')
        job.save()
        preview(job)
        set_preview_sample(job.pk, self.editor, 2, False, True)
        set_preview_sample(job.pk, self.editor, 3, False, True)
        job.refresh_from_db()
        self.assertFalse(job.result['rows'][0]['after']['is_sample'])
        self.assertFalse(job.result['rows'][1]['after']['is_sample'])
        confirm(job.pk, self.editor)
        part.refresh_from_db()
        self.assertFalse(part.is_sample)
        self.assertFalse(Part.objects.get(part_number='NEW-SAMPLE').is_sample)

    def test_preview_classification_rejects_stale_or_inapplicable_rows(self):
        part = create_part({'customer': self.customer, 'part_number': 'EXISTING', 'wall': 1}, self.editor)
        job = self.make_job([['LENNOX 1', 'EXISTING', 1]])
        preview(job)
        set_preview_sample(job.pk, self.editor, 2, True, False)
        with self.assertRaises(ValidationError):
            set_preview_sample(job.pk, self.editor, 2, False, False)
        update_part(part.pk, {'wall': 2}, 1, self.editor)
        with self.assertRaises(ValidationError):
            set_preview_sample(job.pk, self.editor, 2, False, True)
        job.refresh_from_db()
        self.assertEqual(job.sample_overrides, {'2': True})
        with self.assertRaises(ValidationError):
            confirm(job.pk, self.editor)
        job.mode = 'add'
        preview(job)
        with self.assertRaises(ValidationError):
            set_preview_sample(job.pk, self.editor, 2, False, True)
        duplicate = self.make_job([['LENNOX 1', 'DUP', 1], ['LENNOX 1', 'DUP', 2]])
        preview(duplicate)
        with self.assertRaises(ValidationError):
            set_preview_sample(duplicate.pk, self.editor, 2, True, False)

    def test_preview_classification_permissions_and_confirmed_job(self):
        self.client.force_login(self.editor)
        job = self.make_job([['LENNOX 1', 'NEW', 1]])
        preview(job)
        url = reverse('catalog:import_preview', args=[job.pk])
        payload = {'action': 'set_sample', 'row': 2, 'is_sample': 'on', 'expected_sample': 'false'}
        other = User.objects.create_user('other-sample-editor', password='test-password')
        other.groups.add(Group.objects.get(name='Editor'))
        self.client.force_login(other)
        self.assertEqual(self.client.post(url, payload).status_code, 404)
        self.client.force_login(self.reader)
        self.assertEqual(self.client.post(url, payload).status_code, 403)
        self.client.force_login(self.editor)
        self.assertEqual(self.client.post(url, {**payload, 'row': 'invalid'}).status_code, 400)
        confirm(job.pk, self.editor)
        response = self.client.post(url, payload)
        self.assertContains(response, 'Solo puede clasificar piezas en una vista previa pendiente.')
        self.assertFalse(Part.objects.get(part_number='NEW').is_sample)

    def test_integration_sample_controls_permissions_and_version(self):
        admin = User.objects.create_superuser('sample-admin', password='test-password')
        self.client.force_login(admin)
        url = reverse('catalog:integrations')
        self.client.post(url, {'action': 'create', 'name': 'No samples'})
        restricted = Integration.objects.get(name='No samples')
        self.assertFalse(restricted.include_samples)
        self.client.post(url, {'action': 'create', 'name': 'With samples', 'include_samples': 'on'})
        allowed = Integration.objects.get(name='With samples')
        self.assertTrue(allowed.include_samples)
        self.assertContains(self.client.get(url), 'Permitir muestras')
        payload = {'action': 'sample_policy', 'pk': allowed.pk, 'policy_version': 1}
        self.assertEqual(self.client.post(url, payload).status_code, 302)
        allowed.refresh_from_db()
        self.assertFalse(allowed.include_samples)
        self.assertEqual(allowed.policy_version, 2)
        # A stale page cannot overwrite the saved choice.
        self.client.post(url, {**payload, 'include_samples': 'on'})
        allowed.refresh_from_db()
        self.assertFalse(allowed.include_samples)
        self.client.post(url, {**payload, 'policy_version': 2})
        allowed.refresh_from_db()
        self.assertEqual(allowed.policy_version, 2)
        self.client.force_login(self.editor)
        self.assertEqual(self.client.post(url, {**payload, 'policy_version': 2, 'include_samples': 'on'}).status_code, 403)
        allowed.refresh_from_db()
        self.assertFalse(allowed.include_samples)

    def test_integration_restriction_enforced_on_all_part_routes(self):
        sample = create_part({'customer': self.customer, 'part_number': 'PRIVATE-SAMPLE', 'is_sample': True}, self.editor)
        normal = create_part({'customer': self.customer, 'part_number': 'NORMAL'}, self.editor)
        _, token = Integration.create_key('Restricted', include_samples=False)
        headers = {'HTTP_AUTHORIZATION': f'Bearer {token}'}
        self.assertEqual(self.client.get('/api/v1/parts/', **headers).json()['count'], 1)
        self.assertEqual(self.client.get('/api/v1/parts/?is_sample=true', **headers).json()['count'], 0)
        self.assertEqual(self.client.get('/api/v1/parts/exact/?code=PRIVATE-SAMPLE', **headers).json()['count'], 0)
        self.assertEqual(self.client.get(f'/api/v1/parts/{sample.pk}/', **headers).status_code, 404)
        self.assertEqual(self.client.get(f'/api/v1/parts/{normal.pk}/', **headers).status_code, 200)
        events = self.client.get('/api/v1/changes/', **headers).json()['results']
        self.assertTrue(events[0]['excluded'])
        self.assertEqual(events[0]['operation'], 'deactivate')
        self.assertNotIn('part_number', events[0]['part'])
        self.assertFalse(events[1]['excluded'])
        # Historical normal snapshots must not reveal a piece now classified as sample.
        update_part(normal.pk, {'is_sample': True}, 1, self.editor)
        events = self.client.get('/api/v1/changes/', **headers).json()['results']
        self.assertTrue(all(event['excluded'] for event in events))
        update_part(normal.pk, {'is_sample': False}, 2, self.editor)
        self.assertFalse(self.client.get('/api/v1/changes/', **headers).json()['results'][-1]['excluded'])

    def test_sample_policy_change_invalidates_cursors_and_bootstrap_pages(self):
        from django.core import signing
        integration, token = Integration.create_key('Policy changes')
        headers = {'HTTP_AUTHORIZATION': f'Bearer {token}'}
        first = self.client.get('/api/v1/bootstrap/', **headers).json()
        self.assertTrue(first['include_samples'])
        integration.include_samples = False
        integration.policy_version += 1
        integration.save()
        response = self.client.get('/api/v1/changes/', {'cursor': first['changes_cursor']}, **headers)
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()['code'], 'resync_required')
        self.assertEqual(self.client.get(first['parts_url'], **headers).status_code, 409)
        legacy = signing.dumps(0, salt='catalog-changes')
        self.assertEqual(self.client.get('/api/v1/changes/', {'cursor': legacy}, **headers).status_code, 409)
        second = self.client.get('/api/v1/bootstrap/', **headers).json()
        self.assertFalse(second['include_samples'])
        self.assertEqual(self.client.get('/api/v1/changes/', {'cursor': second['changes_cursor']}, **headers).status_code, 200)
        _, other_token = Integration.create_key('Other app')
        self.assertEqual(self.client.get('/api/v1/changes/', {'cursor': second['changes_cursor']}, HTTP_AUTHORIZATION=f'Bearer {other_token}').status_code, 400)

    def test_consumer_rebuilds_cache_when_server_sample_permission_changes(self):
        import json
        import runpy
        import sqlite3
        import urllib.error
        from pathlib import Path
        from urllib.parse import urlsplit

        sample = create_part({'customer': self.customer, 'part_number': 'CENTRAL-SAMPLE', 'is_sample': True}, self.editor)
        normal = create_part({'customer': self.customer, 'part_number': 'CENTRAL-NORMAL'}, self.editor)
        integration, token = Integration.create_key('Live cache test')

        def response(request, **kwargs):
            parsed = urlsplit(request.full_url)
            result = self.client.get(parsed.path + ('?' + parsed.query if parsed.query else ''), HTTP_AUTHORIZATION=f'Bearer {token}')
            if result.status_code >= 400:
                raise urllib.error.HTTPError(request.full_url, result.status_code, 'API error', {}, BytesIO(result.content))
            return BytesIO(result.content)

        with TemporaryDirectory() as directory:
            cache = str(Path(directory) / 'cache.sqlite3')
            env = {'UR_BASE_URL': 'http://testserver', 'UR_API_KEY': token, 'UR_CACHE_DB': cache, 'UR_INCLUDE_SAMPLES': 'true'}

            def sync():
                with patch.dict('os.environ', env), patch('urllib.request.urlopen', side_effect=response):
                    result = runpy.run_path(str(Path(__file__).resolve().parent.parent / 'scripts/sample_consumer.py'))
                    result['db'].close()
                with sqlite3.connect(cache) as db:
                    return {row[0] for row in db.execute('SELECT central_uuid FROM parts')}

            self.assertEqual(sync(), {str(sample.pk), str(normal.pk)})
            integration.include_samples = False
            integration.policy_version += 1
            integration.save()
            self.assertEqual(sync(), {str(normal.pk)})
            integration.include_samples = True
            integration.policy_version += 1
            integration.save()
            self.assertEqual(sync(), {str(sample.pk), str(normal.pk)})
            update_part(normal.pk, {'is_sample': True}, 1, self.editor)
            integration.include_samples = False
            integration.policy_version += 1
            integration.save()
            self.assertEqual(sync(), set())
