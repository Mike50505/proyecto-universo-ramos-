import csv
import io
import hashlib
import json
import unicodedata
from collections import Counter
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from .models import Box, Customer, Destination, ImportJob, Part, SheetColor, normalize_code
from .services import create_part, lock_change_stream, update_part

FIELDS = ['customer', 'destination', 'part_number', 'diameter', 'wall', 'weight', 'weight_unit', 'box', 'std', 'sheet_color', 'reference_customer', 'reference_plant', 'is_sample']
LABELS = {'customer': 'CLIENTE', 'destination': 'DESTINO', 'part_number': 'NÚMERO DE PARTE', 'diameter': 'DIÁMETRO', 'wall': 'PARED', 'weight': 'PESO', 'weight_unit': 'UNIDAD DEL PESO (kg/g)', 'box': 'CAJA', 'std': 'STD.', 'sheet_color': 'HOJA COLOR', 'reference_customer': 'CLIENTE (segunda columna)', 'reference_plant': 'PLANTA'}
LABELS['wall_note'] = 'PARED (texto original)'
LABELS['is_sample'] = 'PENDIENTE DE APROBACIÓN'
CATALOGS = {'customer': Customer, 'destination': Destination, 'box': Box, 'sheet_color': SheetColor}

def fold(value):
    text = ''.join(ch for ch in unicodedata.normalize('NFKD', str(value or '').upper()) if not unicodedata.combining(ch))
    return text.strip().replace('.', '').replace('�', 'U')

def workbook(job):
    job.file.open('rb')
    try:
        return load_workbook(io.BytesIO(job.file.read()), read_only=True, data_only=False, keep_links=False)
    finally:
        job.file.close()

def headers(job, sheet):
    wb = workbook(job)
    if sheet not in wb.sheetnames:
        raise ValidationError('La hoja seleccionada no existe.')
    ws = wb[sheet]
    for row_number, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row or 20, 20)), 1):
        names = [str(cell.value or '').strip() for cell in row]
        if any(fold(x) in ('NUMERO DE PARTE', 'NUMERO PARTE') for x in names):
            return row_number, names
    raise ValidationError('No se encontró una fila de encabezados en las primeras 20 filas.')

def auto_mapping(names):
    mapping = {}
    clients = [i + 1 for i, name in enumerate(names) if fold(name) == 'CLIENTE']
    if len(clients) in (1, 2):
        mapping['customer'] = clients[0]
    if len(clients) == 2:
        mapping['reference_customer'] = clients[1]
    for field in FIELDS:
        if field in ('customer', 'reference_customer'):
            continue
        accepted_labels = {fold(LABELS[field])}
        if field == 'is_sample':
            accepted_labels.add('MUESTRA')  # Keep existing workbooks compatible.
        matches = [i + 1 for i, name in enumerate(names) if fold(name) in accepted_labels]
        if len(matches) == 1:
            mapping[field] = matches[0]
    return mapping

def _value(cell):
    if cell.data_type == 'f':
        raise ValidationError('La fila contiene una fórmula; sustituya por un valor literal.')
    return '' if cell.value is None else str(cell.value).strip()

def parse_rows(job):
    wb = workbook(job)
    header_row, _ = headers(job, job.sheet)
    try:
        mapping = {key: int(value) for key, value in job.mapping.items() if key in FIELDS and value}
    except (TypeError, ValueError):
        raise ValidationError('El mapeo de columnas no es válido.')
    if not mapping.get('part_number'):
        raise ValidationError('Mapee Número de parte.')
    if len(set(mapping.values())) != len(mapping) or any(pos < 1 or pos > 100 for pos in mapping.values()):
        raise ValidationError('Cada campo debe corresponder a una columna distinta y válida.')
    ws = wb[job.sheet]
    if (ws.max_column or 0) > 100 or (ws.max_row or 0) > 10050:
        raise ValidationError('El archivo supera 100 columnas o 10,000 filas de datos.')
    result = []
    for row_number, cells in enumerate(ws.iter_rows(min_row=header_row + 1), header_row + 1):
        if not any(cell.value is not None for cell in cells):
            continue
        raw = {}
        for field in FIELDS:
            if field not in mapping:
                continue
            pos = mapping[field]
            cell = cells[pos - 1] if pos <= len(cells) else None
            try:
                raw[field] = _value(cell) if cell else ''
            except ValidationError as exc:
                raw['__error__'] = '; '.join(exc.messages)
                raw[field] = ''
            if field == 'part_number' and cell and cell.data_type == 'n' and cell.value is not None:
                raw['__error__'] = 'El número de parte es numérico en Excel; convierta la celda a texto para preservar ceros iniciales.'
        result.append((row_number, raw))
        if len(result) > 10000:
            raise ValidationError('El límite es 10,000 filas por archivo.')
    return result

def _validate_field(field, value):
    if field == 'weight_unit':
        if not value:
            return ''
        unit = value.strip().lower()
        if unit not in ('kg', 'g'):
            raise ValidationError('UNIDAD DEL PESO debe ser kg o g.')
        return unit
    if field == 'is_sample':
        if fold(value) in ('SI', 'TRUE', '1', 'X'):
            return True
        if fold(value) in ('', 'NO', 'FALSE', '0'):
            return False
        raise ValidationError('PENDIENTE DE APROBACIÓN debe indicar Sí/No, True/False, 1/0 o X.')
    if field in CATALOGS:
        if len(value) > 150:
            raise ValidationError(f'{LABELS[field]} supera 150 caracteres.')
        return value
    if field in ('wall', 'weight'):
        if not value:
            return None
        try:
            decimal = Decimal(value)
        except InvalidOperation:
            if field == 'wall' and len(value) <= 120:
                return None
            raise ValidationError(f'{LABELS[field]} debe ser decimal.')
        if not decimal.is_finite() or decimal < 0 or (decimal and decimal.adjusted() >= 8) or -decimal.as_tuple().exponent > 6:
            raise ValidationError(f'{LABELS[field]} debe ser un decimal no negativo con hasta 6 decimales.')
        return str(decimal)
    max_length = Part._meta.get_field(field).max_length
    if max_length and len(value) > max_length:
        raise ValidationError(f'{LABELS[field]} supera {max_length} caracteres.')
    return value

def evaluate(job):
    rows = parse_rows(job)
    counts = Counter(normalize_code(raw.get('part_number')) for _, raw in rows if raw.get('part_number'))
    output = []
    for number, raw in rows:
        if str(number) in job.sample_overrides:
            raw['is_sample'] = 'Sí' if job.sample_overrides[str(number)] else 'No'
        code = normalize_code(raw.get('part_number'))
        item = {'row': number, 'code': raw.get('part_number', ''), 'raw': {field: raw.get(field, '') for field in FIELDS}, 'status': '', 'reason': '', 'before': {}, 'after': {}, 'catalog_new': [], 'version': None}
        if raw.get('__error__'):
            item.update(status='error', reason=raw['__error__'])
        elif not code:
            item.update(status='error', reason='Número de parte es obligatorio.')
        elif counts[code] > 1:
            item.update(status='duplicate', reason='El código aparece más de una vez en esta hoja; revise las filas antes de importar.')
        else:
            try:
                data = {}
                for field, value in raw.items():
                    if field not in FIELDS:
                        continue
                    clean = _validate_field(field, value)
                    if field in CATALOGS and value:
                        related = CATALOGS[field].objects.filter(name__iexact=value).first()
                        if related and not related.active:
                            raise ValidationError(f'{LABELS[field]} está inactivo: {value}. Revíselo antes de importar.')
                        if not related:
                            item['catalog_new'].append(LABELS[field] + ': ' + value)
                        clean = related.name if related else value
                    data[field] = clean
                    if field == 'wall':
                        data['wall_note'] = value if clean is None and value else ''
                existing = Part.objects.select_related(*CATALOGS.keys()).filter(normalized_number=code).first()
                sample_value = data.get('is_sample', False)
                if existing and (not raw.get('is_sample') and not job.clear_blanks or 'is_sample' not in data):
                    sample_value = existing.is_sample
                item['is_sample'] = sample_value
                if existing:
                    item['version'] = existing.version
                    changed, before = {}, {}
                    for field, value in data.items():
                        if not raw.get('wall' if field == 'wall_note' else field) and not job.clear_blanks:
                            continue
                        old = getattr(existing, field)
                        old_value = old if field == 'is_sample' else str(old) if old is not None else None
                        new_value = value if value != '' or field not in CATALOGS else None
                        if field == 'customer' and not new_value:
                            new_value = 'PENDIENTE'
                        if field in ('wall', 'weight') and old is not None and new_value is not None and Decimal(new_value) == old:
                            continue
                        if old_value != new_value:
                            before[field] = old_value
                            changed[field] = new_value
                    item['before'], item['after'] = before, changed
                    item['status'] = 'update' if changed else 'unchanged'
                    if changed and job.mode == 'add':
                        item['status'] = 'skipped'
                else:
                    if not data.get('customer'):
                        data['customer'] = 'PENDIENTE'
                        item['raw']['customer'] = 'PENDIENTE'
                        if not Customer.objects.filter(name__iexact='PENDIENTE').exists():
                            item['catalog_new'].append(LABELS['customer'] + ': PENDIENTE')
                    item['after'] = data
                    item['status'] = 'new' if job.mode != 'update' else 'skipped'
            except ValidationError as exc:
                item.update(status='error', reason='; '.join(exc.messages))
        output.append(item)
    return output

def preview(job):
    rows = evaluate(job)
    job.result = {'rows': rows, 'counts': dict(Counter(row['status'] for row in rows))}
    job.status = 'preview'
    job.save(update_fields=['result', 'status'])
    return rows

def materialize(values):
    data = {}
    for field, value in values.items():
        if field in CATALOGS:
            if value:
                related = CATALOGS[field].objects.filter(name__iexact=value).first()
                if related is None:
                    related = CATALOGS[field].objects.create(name=value)
                if not related.active:
                    raise ValidationError(f'{LABELS[field]} está inactivo: {value}.')
                data[field] = related
            else:
                data[field] = None
        elif field in ('wall', 'weight'):
            data[field] = Decimal(value) if value is not None else None
        else:
            data[field] = value
    return data


def comparable_rows(rows):
    return [{**row, 'catalog_new': sorted(row['catalog_new'])} for row in rows]


def preview_revision(job):
    return hashlib.sha256(json.dumps(job.result, sort_keys=True).encode()).hexdigest()


def can_classify(job, row):
    return row['status'] in ('new', 'update', 'unchanged') and (job.mode != 'add' or row['status'] == 'new')


@transaction.atomic
def set_all_preview_samples(job_id, actor, is_sample, expected_revision):
    lock_change_stream()
    job = ImportJob.objects.select_for_update().get(pk=job_id, actor=actor)
    if job.status != 'preview':
        raise ValidationError('Solo puede clasificar piezas en una vista previa pendiente.')
    if preview_revision(job) != expected_revision:
        raise ValidationError('La vista previa cambió. Recargue antes de guardar.')
    rows = job.result.get('rows', [])
    if comparable_rows(rows) != comparable_rows(evaluate(job)):
        raise ValidationError('El catálogo o el archivo cambió. Genere otra vista previa.')
    for row in rows:
        if can_classify(job, row):
            job.sample_overrides[str(row['row'])] = is_sample
    job.save(update_fields=['sample_overrides'])
    preview(job)
    return job


@transaction.atomic
def set_preview_sample(job_id, actor, row_number, is_sample, expected_sample):
    lock_change_stream()
    job = ImportJob.objects.select_for_update().get(pk=job_id, actor=actor)
    if job.status != 'preview':
        raise ValidationError('Solo puede clasificar piezas en una vista previa pendiente.')
    rows = job.result.get('rows', [])
    row = next((row for row in rows if row['row'] == row_number), None)
    if row is None or row['status'] not in ('new', 'update', 'unchanged') or (job.mode == 'add' and row['status'] != 'new'):
        raise ValidationError('Esta fila no se puede clasificar con el modo de importación actual. Revise sus errores o cambie el modo.')
    if row.get('is_sample', False) != expected_sample:
        raise ValidationError('La clasificación cambió desde que abrió esta página. Recargue antes de guardar.')
    if comparable_rows(rows) != comparable_rows(evaluate(job)):
        raise ValidationError('El catálogo o el archivo cambió después de la vista previa. Genere otra vista previa.')
    job.sample_overrides[str(row_number)] = is_sample
    job.save(update_fields=['sample_overrides'])
    preview(job)
    return job

@transaction.atomic
def confirm(job_id, actor):
    lock_change_stream()
    job = ImportJob.objects.select_for_update().get(pk=job_id)
    if job.status == 'confirmed':
        return job
    if job.status != 'preview' or job.actor_id != actor.pk:
        raise ValidationError('La vista previa no está disponible para confirmar.')
    old = job.result['rows']
    new = evaluate(job)
    # Earlier previews may have recorded newly created catalogs in JSON key order.
    # Their order has no meaning, but PostgreSQL can reorder mapping keys.
    if comparable_rows(old) != comparable_rows(new):
        raise ValidationError('El catálogo o el archivo cambió después de la vista previa. Genere otra vista previa.')
    if any(row['status'] in ('duplicate', 'error') for row in new):
        raise ValidationError('Resuelva los duplicados y errores antes de confirmar.')
    for row in new:
        if row['status'] == 'new':
            create_part(materialize(row['after']), actor, 'import')
        elif row['status'] == 'update':
            part = Part.objects.get(normalized_number=normalize_code(row['code']))
            update_part(part.pk, materialize(row['after']), row['version'], actor, 'import')
    job.status = 'confirmed'
    job.confirmed_at = timezone.now()
    job.result = {'rows': new, 'counts': dict(Counter(row['status'] for row in new))}
    job.save(update_fields=['status', 'confirmed_at', 'result'])
    return job

def template_bytes():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Piezas'
    ws.append([LABELS[field] if field != 'reference_customer' else 'CLIENTE' for field in FIELDS])
    ws.freeze_panes = 'A2'
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(16, len(str(col[0].value)) + 3)
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()

def error_csv(job):
    stream = io.StringIO()
    writer = csv.writer(stream)
    writer.writerow(['Fila', 'Número de parte', 'Estado', 'Motivo'])
    for row in job.result.get('rows', []):
        if row['status'] in ('duplicate', 'error'):
            writer.writerow([row['row'], row['code'], row['status'], row['reason']])
    return '\ufeff' + stream.getvalue()
